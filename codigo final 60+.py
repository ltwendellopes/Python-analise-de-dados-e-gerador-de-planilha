import os
import glob
import logging
import json
import datetime as dt

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# =============================================================
#  CONFIGURACAO - edite aqui
# =============================================================
CONFIG = {
    "pasta_entrada":  r"C:\Users\wendel.teodoro\OneDrive - PRIME CONSULTORIA E ASSESSORIA EMPRESARIAL LTDA - EPP\Documentos\Codigo 60+\Nova pasta",
    "carteiras_file": r"C:\Users\wendel.teodoro\OneDrive - PRIME CONSULTORIA E ASSESSORIA EMPRESARIAL LTDA - EPP\Documentos\Codigo 60+\Nova pasta\Carteiras Gerenciadoras.xlsx",
    "aprovadas_file": r"C:\Users\wendel.teodoro\OneDrive - PRIME CONSULTORIA E ASSESSORIA EMPRESARIAL LTDA - EPP\Documentos\Codigo 60+\Nova pasta\PRIME - APROVADAS SEM MOVIMENTAÇÃO 01-06-2026.xlsx",
    "pasta_saida":   r"C:\Users\wendel.teodoro\OneDrive - PRIME CONSULTORIA E ASSESSORIA EMPRESARIAL LTDA - EPP\Documentos\Codigo 60+\Nova pasta\relatorios_os",
}
# =============================================================

logging.basicConfig(level=logging.INFO, format="%(asctime)s  %(message)s", datefmt="%H:%M:%S")
log = logging.getLogger(__name__)

THIN   = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
C = {
    "dark_blue":  "1F3864", "mid_blue":   "2E75B6", "light_blue": "BDD7EE",
    "header_txt": "FFFFFF", "alt":        "EBF0FA", "green":      "C6EFCE",
    "red":        "FFC7CE", "saida_bg":   "FF0000", "entrada_bg": "00B050",
    "gray_tot":   "D9D9D9",
}

def fill(color):                    return PatternFill("solid", fgColor=color)
def bfont(color="000000", size=10): return Font(name="Arial", bold=True, color=color, size=size)


# =============================================================
#  SELECAO DO ARQUIVO DE OS
# =============================================================

def escolher_arquivo():
    padrao   = os.path.join(CONFIG["pasta_entrada"], "*.xlsx")
    arquivos = sorted(glob.glob(padrao), key=os.path.getmtime, reverse=True)
    arquivos = [a for a in arquivos
                if not os.path.basename(a).startswith("OS_")
                and "Carteiras"  not in os.path.basename(a)
                and "APROVADAS"  not in os.path.basename(a)]

    if not arquivos:
        print("\nNenhum arquivo .xlsx de OS encontrado em:")
        print(CONFIG["pasta_entrada"])
        return None

    if len(arquivos) == 1:
        print(f"\nArquivo OS encontrado: {os.path.basename(arquivos[0])}")
        return arquivos[0]

    print("\nArquivos encontrados (mais recente primeiro):")
    for i, a in enumerate(arquivos, 1):
        mtime = dt.datetime.fromtimestamp(os.path.getmtime(a)).strftime("%d/%m/%Y %H:%M")
        print(f"  [{i}] {os.path.basename(a)}  ({mtime})")

    while True:
        escolha = input("\nDigite o numero do arquivo (Enter = mais recente): ").strip()
        if escolha == "":
            return arquivos[0]
        if escolha.isdigit() and 1 <= int(escolha) <= len(arquivos):
            return arquivos[int(escolha) - 1]
        print("Opcao invalida.")


# =============================================================
#  CONSTANTES — OS
# =============================================================

STATUS_INDESEJADOS        = ["APROVADA", "CANCELADA", "SERVICO REJEITADO"]
AUTOGESTAO                = ["Auto Gestao", "Autogestao", "Auto Gestão", "Autogestão"]
ATENDIMENTOS_NF           = ["Equipe Especializada", "Validacao de NF", "Validação de NF",
                              "Equipe Orcamentista", "Equipe Orçamentista"]
ABAS_CLIENTE              = [
    "Aguard. Aprovacao", "Aguard. Aprovação",
    "Aguard. Envio Veiculo", "Aguard. Envio Veículo",
    "NF Validadas",
    "Saldo Insuf. Lib. Aprovacao", "Saldo Insuf. Lib. Aprovação",
    "Fim Servico", "Fim Serviço",
    "Inicio Servico", "Inicio Serviço",
]
ABAS_ORCAMENTISTA_CLIENTE = [
    "Liberacao Aprovacao", "Liberação Aprovação",
    "NF Validadas",
    "Saldo Insuf. Lib. Aprovacao", "Saldo Insuf. Lib. Aprovação",
    "Validacao 1 Orcamento", "Validação 1º Orçamento",
    "OS Plataforma",
]
ABAS_SO_NF = [
    "OS Plataforma", "Aguard. 1 Orcamento", "Aguard. 1º Orçamento",
    "Aguard. Reavaliacao", "Aguard. Reavaliação", "Cotacoes", "Cotações",
    "Liberacao Aprovacao", "Liberação Aprovação", "NF Validadas",
    "Saldo Insuf. Lib. Aprovacao", "Saldo Insuf. Lib. Aprovação",
    "Validacao 1 Orcamento", "Validação 1º Orçamento",
    "Aguard. Envio Veiculo", "Aguard. Envio Veículo",
]


# =============================================================
#  TRATAMENTO — OS
# =============================================================

def tratar_dados_os(caminho_arquivo):
    tabela          = pd.read_excel(caminho_arquivo, sheet_name=0)
    tabela_carteira = pd.read_excel(CONFIG["carteiras_file"], sheet_name=0)

    tabela.columns          = tabela.columns.str.strip()
    tabela_carteira.columns = tabela_carteira.columns.str.strip()

    for col in ["Validação NF", "Validacao NF", "Atendimento", "Gerenciadora"]:
        if col in tabela.columns:
            tabela.drop(columns=col, inplace=True)

    cart_cols = ["Cod", "Lider", "Status", "Atendimento"]
    for opt in ["Validação NF", "Validacao NF"]:
        if opt in tabela_carteira.columns:
            cart_cols.append(opt)
            break
    if "Gerenciadora" in tabela_carteira.columns:
        cart_cols.append("Gerenciadora")

    tabela_final = pd.merge(
        tabela, tabela_carteira[cart_cols],
        left_on="Cod Cliente", right_on="Cod", how="left",
    )
    tabela_final.drop(columns="Cod", inplace=True)
    tabela_final["cliente_os"] = (
        tabela_final["Cod Cliente"].astype(str) + "-" + tabela_final["Cod OS"].astype(str)
    )

    tabela_final = tabela_final[~tabela_final["Status OS"].isin(STATUS_INDESEJADOS)]
    tabela_final = tabela_final.drop_duplicates(subset="cliente_os")
    tabela_final = tabela_final[tabela_final["Qtd Dias em Atraso"] >= 60]
    tabela_final = tabela_final[~tabela_final["Atendimento"].isin(AUTOGESTAO)]
    tabela_final = tabela_final[~tabela_final["Status"].isin(["Inativo"])]

    eh_prime = tabela_final["Gerenciadora"].str.strip().str.upper() == "PRIME"

    tabela_final.loc[
        eh_prime &
        (tabela_final["Lider"] == "Anderson de Oliveira") &
        (tabela_final["Atendimento"].str.strip() == "Equipe Especializada"),
        "Lider"] = "BOP"

    vnf_col = "Validação NF" if "Validação NF" in tabela_final.columns else "Validacao NF"
    tabela_final["_vnf"] = tabela_final[vnf_col].str.strip().str.upper()

    tabela_final.loc[
        eh_prime &
        (tabela_final["Aba OS"] == "Validação NF") &
        (tabela_final["_vnf"] == "EQUIPE PLATAFORMA") &
        (tabela_final["Atendimento"].str.strip().isin(ATENDIMENTOS_NF)),
        "Lider"] = "NF"
    tabela_final.loc[
        eh_prime &
        (tabela_final["Aba OS"] == "NF c/ Problema") &
        (tabela_final["_vnf"] == "EQUIPE PLATAFORMA") &
        (tabela_final["Atendimento"].str.strip().isin(ATENDIMENTOS_NF)),
        "Lider"] = "NF/Problema"

    # Para TODAS as gerenciadoras (NF Link, NEO, etc.):
    # Se VNF não é "EQUIPE PLATAFORMA" nessas abas → Lider = Cliente
    mask_vnf_nao_plat = (
        tabela_final["Aba OS"].isin(["Validação NF", "NF c/ Problema"]) &
        (tabela_final["_vnf"] != "EQUIPE PLATAFORMA")
    )
    tabela_final.loc[mask_vnf_nao_plat, "Lider"] = "Cliente"

    tabela_final.drop(columns="_vnf", inplace=True)


    mask_orc = (
        tabela_final["Aba OS"].isin(ABAS_ORCAMENTISTA_CLIENTE) &
        (tabela_final["Atendimento"].str.strip().isin(["Equipe Orcamentista", "Equipe Orçamentista"]))
    )
    tabela_final.loc[mask_orc, "Lider"] = "Cliente"

    vnf_col2 = "Validação NF" if "Validação NF" in tabela_final.columns else "Validacao NF"
    mask_vnf_orc = (
        (tabela_final["Aba OS"] == "Validação NF") &
        (tabela_final["Atendimento"].str.strip().isin(["Equipe Orcamentista", "Equipe Orçamentista"])) &
        (tabela_final[vnf_col2].str.strip().str.upper() == "CLIENTE")
    )
    tabela_final.loc[mask_vnf_orc, "Lider"] = "Cliente"
    tabela_final.loc[tabela_final["Aba OS"].isin(ABAS_CLIENTE), "Lider"] = "Cliente"

    vnf_col = "Validação NF" if "Validação NF" in tabela_final.columns else "Validacao NF"
    tabela_final = tabela_final[~(
        (tabela_final[vnf_col] == "Cliente") &
        (tabela_final["Aba OS"].isin(["NF c/ Problema", "Validação NF"]))
    )]
    tabela_final = tabela_final[~(
        (tabela_final["Atendimento"].fillna("").str.strip().isin(
            ["Validacao de NF", "Validação de NF"])) &
        (tabela_final["Aba OS"].isin(ABAS_SO_NF))
    )]

    log.info(f"[OS] Total final: {len(tabela_final)} linhas")
    return tabela_final


# =============================================================
#  TRATAMENTO — APROVADAS
#  Colunas da planilha:
#  Gerenciadora | Data de Abertura | Data de Aprovação | Código Cliente |
#  Código OS | Placa do Veículo | CNPJ | Razão Social |
#  Valor Total de Itens | Valor Total de MDO | Checklist | NF
# =============================================================

# Clientes que ficam com o lider original na PRIME (não viram Jaqueline)
COD_EXCECAO_JAQUELINE = [10312]

def tratar_dados_aprovadas():
    """
    Lê APROVADAS.xlsx (que já tem coluna Gerenciadora própria).
    Mescla com Carteiras apenas para obter o Lider.

    Regra PRIME:
      - Todos os registros da PRIME viram "Jaqueline Barbosa",
        EXCETO o cliente 10312, que mantém o líder original da carteira.

    Outras gerenciadoras: lider original, sem alteração.
    """
    if not os.path.exists(CONFIG["aprovadas_file"]):
        log.warning("Arquivo APROVADAS nao encontrado. Pulando.")
        return pd.DataFrame()

    tabela          = pd.read_excel(CONFIG["aprovadas_file"], sheet_name=0)
    tabela_carteira = pd.read_excel(CONFIG["carteiras_file"], sheet_name=0)

    tabela.columns          = tabela.columns.str.strip()
    tabela_carteira.columns = tabela_carteira.columns.str.strip()

    # Da carteira só precisamos do Lider (Gerenciadora já vem da planilha APROVADAS)
    cart_cols = ["Cod", "Lider"]
    tabela_final = pd.merge(
        tabela,
        tabela_carteira[cart_cols],
        left_on="Código Cliente", right_on="Cod",
        how="left",
    )
    tabela_final.drop(columns="Cod", inplace=True)

    # Chave única para deduplicação
    tabela_final["cliente_os"] = (
        tabela_final["Código Cliente"].astype(str) + "-" +
        tabela_final["Código OS"].astype(str)
    )
    tabela_final = tabela_final.drop_duplicates(subset="cliente_os")

    # Receita Parada = soma dos dois valores
    tabela_final["Receita Parada"] = (
        tabela_final["Valor Total de Itens"].fillna(0) +
        tabela_final["Valor Total de MDO"].fillna(0)
    )

    # Dias em aberto desde a aprovação
    hoje = pd.to_datetime(dt.datetime.today().date())
    tabela_final["Dias em aberto"] = (
        hoje - pd.to_datetime(tabela_final["Data de Aprovação"])
    ).dt.days

    # Filtro 60+
    tabela_final = tabela_final[tabela_final["Dias em aberto"] >= 60].copy()

    # ── Regra PRIME ───────────────────────────────────────────────
    # Todos os registros da PRIME viram Jaqueline Barbosa,
    # exceto o cliente 10312 que mantém o lider original.
    eh_prime     = tabela_final["Gerenciadora"].str.strip().str.upper() == "PRIME"
    eh_excecao   = tabela_final["Código Cliente"].isin(COD_EXCECAO_JAQUELINE)
    mask_jaqueline = eh_prime & ~eh_excecao
    tabela_final.loc[mask_jaqueline, "Lider"] = "Jaqueline Barbosa"
    # ─────────────────────────────────────────────────────────────

    log.info(f"[APROVADAS] Total 60+: {len(tabela_final)} linhas")
    log.info(f"[APROVADAS] Gerenciadoras: {sorted(tabela_final['Gerenciadora'].dropna().unique().tolist())}")
    return tabela_final


# =============================================================
#  HISTORICO — separado por gerenciadora e por fonte
# =============================================================

def _path_historico(ger, fonte):
    nome = ger.replace("/", "-").replace(" ", "_")
    return os.path.join(CONFIG["pasta_saida"], f"historico_{fonte}_{nome}.json")

def carregar_historico(ger, fonte):
    p = _path_historico(ger, fonte)
    if os.path.exists(p):
        with open(p, encoding="utf-8") as f:
            return json.load(f)
    return {}

def salvar_historico_os(df_ger, ger, historico):
    os.makedirs(CONFIG["pasta_saida"], exist_ok=True)
    hoje     = dt.datetime.now().strftime("%d/%m/%y")
    clientes = int((df_ger["Lider"] == "Cliente").sum())
    plat     = len(df_ger) - clientes
    historico[hoje] = {
        "plataforma": plat,
        "cliente":    clientes,
        "por_lider":  {k: int(v) for k, v in df_ger.groupby("Lider").size().items()},
    }
    with open(_path_historico(ger, "os"), "w", encoding="utf-8") as f:
        json.dump(historico, f, ensure_ascii=False, indent=2)
    log.info(f"[OS] Historico [{ger}] salvo — {hoje}")

def salvar_historico_aprov(df_ger, ger, historico):
    os.makedirs(CONFIG["pasta_saida"], exist_ok=True)
    hoje = dt.datetime.now().strftime("%d/%m/%y")
    por_lider = {}
    for lider, grp in df_ger.groupby("Lider"):
        por_lider[lider] = {
            "qtd":     int(len(grp)),
            "receita": float(grp["Receita Parada"].sum()),
        }
    historico[hoje] = {
        "total_qtd":     int(len(df_ger)),
        "total_receita": float(df_ger["Receita Parada"].sum()),
        "por_lider":     por_lider,
    }
    with open(_path_historico(ger, "aprov"), "w", encoding="utf-8") as f:
        json.dump(historico, f, ensure_ascii=False, indent=2)
    log.info(f"[APROV] Historico [{ger}] salvo — {hoje}")


# =============================================================
#  HELPERS DE CÉLULA
# =============================================================

def _hdr(ws, r, c, v, bg=None, fg="FFFFFF", size=10):
    cell = ws.cell(row=r, column=c, value=v)
    cell.font      = Font(name="Arial", bold=True, color=fg, size=size)
    cell.fill      = fill(bg or C["dark_blue"])
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border    = BORDER

def _dat(ws, r, c, v, bg=None, bold=False, fmt=None, align="center"):
    cell = ws.cell(row=r, column=c, value=v)
    cell.font      = Font(name="Arial", bold=bold, size=10)
    cell.alignment = Alignment(horizontal=align, vertical="center")
    cell.border    = BORDER
    if bg:  cell.fill          = fill(bg)
    if fmt: cell.number_format = fmt

def _titulo(ws, r, c1, c2, texto, bg=None):
    ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
    cell = ws.cell(row=r, column=c1, value=texto)
    cell.font      = Font(name="Arial", bold=True, color="FFFFFF", size=11)
    cell.fill      = fill(bg or C["mid_blue"])
    cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[r].height = 22

def _mesclar_hdr(ws, r, c1, c2, texto, bg=None):
    ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
    cell = ws.cell(row=r, column=c1, value=texto)
    cell.font      = bfont("FFFFFF")
    cell.fill      = fill(bg or C["mid_blue"])
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border    = BORDER
    for c in range(c1 + 1, c2 + 1):
        ws.cell(row=r, column=c).border = BORDER

def _auto_width(ws):
    for col in ws.columns:
        w = max((len(str(c.value or "")) for c in col), default=10)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(w + 2, 10), 42)


# =============================================================
#  TABELAS — OS
# =============================================================

LIDERES_ORDEM_OS = [
    "Equipe DF", "Vinicius Raimundo Navakoski", "Jose Almy",
    "NF", "NF/Problema", "Diego Rosa", "Paulo videira",
    "BOP", "Anderson de Oliveira", "Rafaela Rocha", "Cliente",
]
ABAS_COLUNAS_OS = [
    "Aguard. 1º Orçamento", "Aguard. Reavaliação", "Cotações",
    "Liberação Aprovação", "NF c/ Problema", "OS Plataforma",
    "Validação 1º Orçamento", "Validação NF",
]

def _ord_os(df, col="Lider"):
    ordem = {l: i for i, l in enumerate(LIDERES_ORDEM_OS)}
    df = df.copy()
    df["_o"] = df[col].map(lambda x: ordem.get(x, 999))
    return df.sort_values("_o").drop(columns="_o")

def _arvore(ws, df, r0):
    lt = df.groupby("Lider").size().reset_index(name="Q")
    la = df.groupby(["Lider", "Aba OS"]).size().reset_index(name="Q")
    _titulo(ws, r0, 1, 3, "Lider / Aba OS / Quantidade")
    for c, h in enumerate(["Lider", "Aba OS", "Qtd"], 1):
        _hdr(ws, r0+1, c, h)
    r = r0 + 2
    for _, lr in _ord_os(lt).iterrows():
        lider = lr["Lider"]
        _dat(ws, r, 1, lider, bg=C["light_blue"], bold=True, align="left")
        _dat(ws, r, 2, "",    bg=C["light_blue"])
        _dat(ws, r, 3, int(lr["Q"]), bg=C["light_blue"], bold=True)
        r += 1
        for _, sr in la[la["Lider"] == lider].sort_values("Aba OS").iterrows():
            _dat(ws, r, 1, "",           bg=C["alt"])
            _dat(ws, r, 2, sr["Aba OS"], bg=C["alt"], align="left")
            _dat(ws, r, 3, int(sr["Q"]), bg=C["alt"])
            r += 1
    _dat(ws, r, 1, "Total Geral", bold=True, bg=C["dark_blue"])
    ws.cell(row=r, column=1).font = bfont("FFFFFF")
    _dat(ws, r, 2, "", bg=C["dark_blue"])
    _dat(ws, r, 3, int(lt["Q"].sum()), bold=True, bg=C["dark_blue"])
    ws.cell(row=r, column=3).font = bfont("FFFFFF")
    return r + 2

def _datas(ws, df_sem_cli, df_cli, hist, cs=5, r0=1):
    ce = cs + len(hist) + 1
    _titulo(ws, r0, cs, ce, "60 Dias Todas Abas (Exceto Aprovadas)")
    for c, h in enumerate(["Data"] + list(hist.keys()) + ["Total +60 geral"], cs):
        _hdr(ws, r0+1, c, h)
    for lbl, row_off in [("Atuacao Plataforma", 2), ("Atuacao Cliente", 3)]:
        cell = ws.cell(row=r0+row_off, column=cs, value=lbl)
        cell.font = bfont(); cell.border = BORDER
    for i, (d, v) in enumerate(hist.items()):
        _dat(ws, r0+2, cs+1+i, v.get("plataforma", ""))
        _dat(ws, r0+3, cs+1+i, v.get("cliente", ""))
    c_hoje = cs + len(hist)
    _dat(ws, r0+2, c_hoje, int(len(df_sem_cli)))
    _dat(ws, r0+3, c_hoje, int(len(df_cli)))
    ws.merge_cells(start_row=r0+2, start_column=ce, end_row=r0+3, end_column=ce)
    cell = ws.cell(row=r0+2, column=ce, value=int(len(df_sem_cli)))
    cell.font = Font(name="Arial", bold=True, size=12)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = BORDER
    return r0 + 5

def _total_os(ws, df_ger, base_ger, cs=5, r0=8):
    ce = cs + 3
    _titulo(ws, r0, cs, ce, "Total - OS")
    for c, h in enumerate(["Totais", "Qtde OS", "Porcentagem", "Total geral"], cs):
        _hdr(ws, r0+1, c, h)
    t60  = len(df_ger)
    t59  = max(0, base_ger - t60)
    base = base_ger
    for i, (lbl, qty, pct) in enumerate([
        ("Total 60+",   t60,  t60  / base if base else 0),
        ("Total 59-",   t59,  t59  / base if base else 0),
        ("Total geral", base, 1.0),
    ]):
        rr = r0 + 2 + i; bold = (i == 2)
        _dat(ws, rr, cs,   lbl,  bold=bold)
        _dat(ws, rr, cs+1, qty,  bold=bold)
        _dat(ws, rr, cs+2, pct,  bold=bold, fmt="0.00%")
        if i == 0: _dat(ws, rr, ce, base, bold=True)
        else:      ws.cell(row=rr, column=ce).border = BORDER
    return r0 + 6

def _lider_data(ws, df, hist, cs=5, r0=15):
    datas = list(hist.keys())
    ce    = cs + len(datas) + 2
    _titulo(ws, r0, cs, ce, "60+ Todas as Abas - Quantidade - Lider x Por Data")
    _mesclar_hdr(ws, r0+1, ce-1, ce,
                 dt.datetime.now().strftime("%d/%m/%Y") + " (Qtde OS)",
                 bg=C["mid_blue"])
    for c, h in enumerate(
        ["Lideres"] + [d + " (Qtde OS)" for d in datas] + ["saidas", "entrada"], cs
    ):
        bg = C["saida_bg"] if h == "saidas" else C["entrada_bg"] if h == "entrada" else C["dark_blue"]
        _hdr(ws, r0+2, c, h, bg=bg)
    lh = _ord_os(df.groupby("Lider").size().reset_index(name="Q"))
    tl = {d: hist[d].get("por_lider", {}) for d in datas}
    r = r0 + 3; ss = se = 0
    for _, lr in lh.iterrows():
        lider = lr["Lider"]; hv = int(lr["Q"]); ant = 0
        _dat(ws, r, cs, lider, align="left")
        for i, d in enumerate(datas):
            v = int(tl[d].get(lider, 0)); _dat(ws, r, cs+1+i, v); ant = v
        s = max(0, ant - hv); e = max(0, hv - ant); ss += s; se += e
        _dat(ws, r, ce-1, s, bg=C["red"]   if s else None)
        _dat(ws, r, ce,   e, bg=C["green"] if e else None)
        r += 1
    _dat(ws, r, cs, "Total", bold=True, bg=C["dark_blue"])
    ws.cell(row=r, column=cs).font = bfont("FFFFFF")
    for i, d in enumerate(datas):
        t = int(sum(tl[d].values()))
        _dat(ws, r, cs+1+i, t, bold=True, bg=C["dark_blue"])
        ws.cell(row=r, column=cs+1+i).font = bfont("FFFFFF")
    total_hoje = int(lh["Q"].sum())
    _dat(ws, r, cs + len(datas), total_hoje, bold=True, bg=C["dark_blue"])
    ws.cell(row=r, column=cs + len(datas)).font = bfont("FFFFFF")
    _dat(ws, r, ce-1, ss, bold=True, bg=C["saida_bg"])
    _dat(ws, r, ce,   se, bold=True, bg=C["entrada_bg"])
    ws.cell(row=r, column=ce-1).font = bfont("FFFFFF")
    ws.cell(row=r, column=ce).font   = bfont("FFFFFF")
    return r + 2

def _lider_aba(ws, df, cs=5, r0=30):
    abas = [a for a in ABAS_COLUNAS_OS if a in df["Aba OS"].unique()]
    ce   = cs + len(abas)
    _titulo(ws, r0, cs, ce, "60+ Todas as Abas - Quantidade - Lider x Por Aba")
    for c, h in enumerate(["Lideres"] + abas, cs):
        _hdr(ws, r0+1, c, h)
    pivot = df.groupby(["Lider", "Aba OS"]).size().unstack(fill_value=0).reset_index()
    for a in abas:
        if a not in pivot.columns: pivot[a] = 0
    r = r0 + 2
    for _, pr in _ord_os(pivot).iterrows():
        _dat(ws, r, cs, pr["Lider"], align="left")
        for i, a in enumerate(abas):
            _dat(ws, r, cs+1+i, int(pr.get(a, 0)))
        r += 1
    _dat(ws, r, cs, "Total", bold=True, bg=C["dark_blue"])
    ws.cell(row=r, column=cs).font = bfont("FFFFFF")
    for i, a in enumerate(abas):
        t = int(pivot[a].sum()) if a in pivot.columns else 0
        _dat(ws, r, cs+1+i, t, bold=True, bg=C["dark_blue"])
        ws.cell(row=r, column=cs+1+i).font = bfont("FFFFFF")
    return r + 2


# =============================================================
#  TABELAS — APROVADAS
# =============================================================

def _aprov_lider_data(ws, df_ger, hist_aprov, cs=5, r0=1):
    """
    Tabela '60 Dias - APROVADAS':
    Líderes | Qtde.OS (ant) | Receita (ant) | Qtde.OS (hoje) | Receita (hoje) |
    % | Saída | Saída receita | Entrada | Entrada receita
    """
    datas     = list(hist_aprov.keys())
    data_hoje = datas[-1] if datas else dt.datetime.now().strftime("%d/%m/%y")
    data_ant  = datas[-2] if len(datas) >= 2 else None

    # índices das colunas
    C_LID  = cs
    C_QANT = cs + 1;  C_RANT = cs + 2
    C_QHOJ = cs + 3;  C_RHOJ = cs + 4;  C_PCT = cs + 5
    C_SAID = cs + 6;  C_SREC = cs + 7
    C_ENT  = cs + 8;  C_EREC = cs + 9
    CE     = C_EREC

    # Linha 1 — título geral
    _titulo(ws, r0, C_LID, CE, "60 Dias - APROVADAS")

    # Linha 2 — sub-títulos de grupos
    # coluna Líderes (vazia no nível de grupo)
    ws.cell(row=r0+1, column=C_LID).border = BORDER

    label_ant = f"Qtde.OS {data_ant}" if data_ant else "Qtde.OS (anterior)"
    _mesclar_hdr(ws, r0+1, C_QANT, C_RANT, label_ant, bg=C["mid_blue"])

    label_hoj = f"Qtde.OS {data_hoje}"
    _mesclar_hdr(ws, r0+1, C_QHOJ, CE, label_hoj, bg=C["mid_blue"])

    # Linha 3 — cabeçalhos de coluna
    hdrs = [
        (C_LID,  "Líderes",          C["dark_blue"]),
        (C_QANT, "Qtde.OS",          C["dark_blue"]),
        (C_RANT, "Receita Parada $",  C["dark_blue"]),
        (C_QHOJ, "Qtde.OS",          C["dark_blue"]),
        (C_RHOJ, "Receita Parada $",  C["dark_blue"]),
        (C_PCT,  "Porcentagem",       C["dark_blue"]),
        (C_SAID, "Saida",             C["saida_bg"]),
        (C_SREC, "Saida receita",     C["saida_bg"]),
        (C_ENT,  "Entrada",           C["entrada_bg"]),
        (C_EREC, "Entrada receita",   C["entrada_bg"]),
    ]
    for col, txt, bg in hdrs:
        _hdr(ws, r0+2, col, txt, bg=bg)

    # Agrega dados atuais da gerenciadora
    grp_hoje = (
        df_ger.groupby("Lider")
        .agg(qtd=("Receita Parada", "count"), rec=("Receita Parada", "sum"))
        .reset_index()
        .sort_values("qtd", ascending=False)
    )

    # Dados históricos da data anterior
    hist_ant = {}
    if data_ant and data_ant in hist_aprov:
        hist_ant = hist_aprov[data_ant].get("por_lider", {})

    total_hoje_qtd = int(grp_hoje["qtd"].sum())
    total_hoje_rec = float(grp_hoje["rec"].sum())
    total_ant_qtd  = int(sum(v.get("qtd", 0)     for v in hist_ant.values()))
    total_ant_rec  = float(sum(v.get("receita", 0) for v in hist_ant.values()))

    r = r0 + 3
    tot_said_q = tot_said_r = tot_ent_q = tot_ent_r = 0.0

    for _, row in grp_hoje.iterrows():
        lider = row["Lider"]
        q_hoj = int(row["qtd"])
        r_hoj = float(row["rec"])
        pct   = q_hoj / total_hoje_qtd if total_hoje_qtd else 0.0

        ant   = hist_ant.get(lider, {})
        q_ant = int(ant.get("qtd", 0))
        r_ant = float(ant.get("receita", 0.0))

        s_q = max(0, q_ant - q_hoj); s_r = max(0.0, r_ant - r_hoj)
        e_q = max(0, q_hoj - q_ant); e_r = max(0.0, r_hoj - r_ant)
        tot_said_q += s_q; tot_said_r += s_r
        tot_ent_q  += e_q; tot_ent_r  += e_r

        _dat(ws, r, C_LID,  lider, align="left")
        _dat(ws, r, C_QANT, q_ant)
        _dat(ws, r, C_RANT, r_ant, fmt='#,##0.00')
        _dat(ws, r, C_QHOJ, q_hoj)
        _dat(ws, r, C_RHOJ, r_hoj, fmt='#,##0.00')
        _dat(ws, r, C_PCT,  pct,   fmt="0.00%")
        _dat(ws, r, C_SAID, s_q,   bg=C["red"]   if s_q else None)
        _dat(ws, r, C_SREC, s_r,   bg=C["red"]   if s_r else None, fmt='#,##0.00')
        _dat(ws, r, C_ENT,  e_q,   bg=C["green"] if e_q else None)
        _dat(ws, r, C_EREC, e_r,   bg=C["green"] if e_r else None, fmt='#,##0.00')
        r += 1

    # Linha Total geral
    _dat(ws, r, C_LID,  "Total geral",      bold=True, bg=C["gray_tot"])
    _dat(ws, r, C_QANT, int(total_ant_qtd), bold=True, bg=C["gray_tot"])
    _dat(ws, r, C_RANT, total_ant_rec,      bold=True, bg=C["gray_tot"], fmt='#,##0.00')
    _dat(ws, r, C_QHOJ, total_hoje_qtd,     bold=True, bg=C["gray_tot"])
    _dat(ws, r, C_RHOJ, total_hoje_rec,     bold=True, bg=C["gray_tot"], fmt='#,##0.00')
    _dat(ws, r, C_PCT,  1.0,                bold=True, bg=C["gray_tot"], fmt="0.00%")
    _dat(ws, r, C_SAID, int(tot_said_q),    bold=True, bg=C["saida_bg"])
    _dat(ws, r, C_SREC, tot_said_r,         bold=True, bg=C["saida_bg"], fmt='#,##0.00')
    _dat(ws, r, C_ENT,  int(tot_ent_q),     bold=True, bg=C["entrada_bg"])
    _dat(ws, r, C_EREC, tot_ent_r,          bold=True, bg=C["entrada_bg"], fmt='#,##0.00')
    for cc in [C_SAID, C_SREC, C_ENT, C_EREC]:
        ws.cell(row=r, column=cc).font = bfont("FFFFFF")

    return r + 2


def _aprov_totais(ws, df_ger, base_aprov_ger, cs=5, r0=1):
    """
    Tabela 'APROVADAS - Total':
    Totais | Qtde.OS (hoje) | Receita Parada $ | Porcentagem
    """
    CE = cs + 3
    _titulo(ws, r0, cs, CE, "APROVADAS - Total")
    hoje_label = dt.datetime.now().strftime("%d/%m")
    for c, h in enumerate(
        ["Totais", f"Qtde.OS {hoje_label}", "Receita Parada $", "Porcentagem"], cs
    ):
        _hdr(ws, r0+1, c, h)

    t60_q  = int(len(df_ger))
    t60_r  = float(df_ger["Receita Parada"].sum())
    base_q = int(base_aprov_ger["qtd"])
    base_r = float(base_aprov_ger["receita"])
    t59_q  = max(0, base_q - t60_q)
    t59_r  = max(0.0, base_r - t60_r)

    for i, (lbl, qty, rec, pct) in enumerate([
        ("Total 60+",   t60_q,  t60_r,  t60_q / base_q if base_q else 0),
        ("Total 59-",   t59_q,  t59_r,  t59_q / base_q if base_q else 0),
        ("Total geral", base_q, base_r, 1.0),
    ]):
        rr = r0 + 2 + i; bold = (i == 2)
        _dat(ws, rr, cs,   lbl,  bold=bold)
        _dat(ws, rr, cs+1, qty,  bold=bold)
        _dat(ws, rr, cs+2, rec,  bold=bold, fmt='#,##0.00')
        _dat(ws, rr, cs+3, pct,  bold=bold, fmt="0.00%")

    return r0 + 6


# =============================================================
#  ABA DE DADOS GENÉRICA
# =============================================================

def _aba_dados(wb, df, titulo):
    ws = wb.create_sheet(title=titulo[:31])
    ws.freeze_panes = "A3"
    ws.sheet_view.showGridLines = False
    ncols = len(df.columns)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    cell = ws.cell(row=1, column=1, value=titulo)
    cell.font      = Font(name="Arial", bold=True, color="FFFFFF", size=12)
    cell.fill      = fill(C["dark_blue"])
    cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 24
    for c, col in enumerate(df.columns, 1):
        _hdr(ws, 2, c, col)
    ws.row_dimensions[2].height = 28
    for r_idx, row in enumerate(df.itertuples(index=False), start=3):
        alt = (r_idx % 2 == 0)
        for c_idx, val in enumerate(row, start=1):
            v = None if (isinstance(val, float) and pd.isna(val)) else val
            cell = ws.cell(row=r_idx, column=c_idx, value=v)
            cell.font      = Font(name="Arial", size=10)
            cell.border    = BORDER
            cell.alignment = Alignment(vertical="center", horizontal="center")
            if alt: cell.fill = fill(C["alt"])
    for c_idx, col in enumerate(df.columns, 1):
        max_len = max(len(str(col)), df[col].astype(str).str.len().max())
        ws.column_dimensions[get_column_letter(c_idx)].width = min(max_len + 3, 40)
    ws.auto_filter.ref = f"A2:{get_column_letter(ncols)}{len(df)+2}"


# =============================================================
#  MONTAGEM DO WORKBOOK
# =============================================================

def build_sheet(wb, ger, df_os_ger, hist_os, base_os_ger,
                df_aprov_ger, hist_aprov, base_aprov_ger):

    df_os_sem_cli = df_os_ger[df_os_ger["Lider"] != "Cliente"]
    df_os_cli     = df_os_ger[df_os_ger["Lider"] == "Cliente"]

    # ── Aba Resumo ───────────────────────────────────────────────
    ws = wb.create_sheet(title="Resumo")
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A3"

    ws.merge_cells("A1:D1")
    cell = ws["A1"]
    cell.value     = f"Gerenciadora: {ger}  |  {dt.datetime.now():%d/%m/%Y %H:%M}"
    cell.font      = Font(name="Arial", bold=True, color="FFFFFF", size=12)
    cell.fill      = fill(C["dark_blue"])
    cell.alignment = Alignment(horizontal="center")
    ws.row_dimensions[1].height = 25

    # Bloco OS — árvore (cols A-D) + tabelas (cols E+)
    _arvore(ws, df_os_sem_cli, r0=3)
    r = _datas(ws,    df_os_sem_cli, df_os_cli, hist_os, cs=5, r0=3)
    r = _total_os(ws, df_os_ger,                base_os_ger,  cs=5, r0=r)
    r = _lider_data(ws, df_os_sem_cli, hist_os, cs=5, r0=r)
    r = _lider_aba(ws,  df_os_sem_cli,          cs=5, r0=r)

    # Bloco APROVADAS — logo abaixo (2 linhas de espaço), mesmas colunas E+
    r_aprov = r + 2
    if not df_aprov_ger.empty:
        r = _aprov_lider_data(ws, df_aprov_ger, hist_aprov, cs=5, r0=r_aprov)
        _aprov_totais(ws, df_aprov_ger, base_aprov_ger, cs=5, r0=r)
    else:
        ws.cell(row=r_aprov, column=5,
                value=f"Sem dados APROVADAS 60+ para {ger}").font = bfont()

    _auto_width(ws)
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 10

    # ── Abas de dados ────────────────────────────────────────────
    # OS plataforma (sem cliente)
    cols_os = [c for c in df_os_sem_cli.columns if c != "cliente_os"]
    _aba_dados(wb, df_os_sem_cli[cols_os].reset_index(drop=True), "Dados OS")

    # OS cliente
    if not df_os_cli.empty:
        cols_cli = [c for c in df_os_cli.columns if c != "cliente_os"]
        _aba_dados(wb, df_os_cli[cols_cli].reset_index(drop=True), "Cliente OS")

    # Todos OS (plataforma + cliente)
    cols_all = [c for c in df_os_ger.columns if c != "cliente_os"]
    _aba_dados(wb, df_os_ger[cols_all].reset_index(drop=True), "Todos OS")

    # Aprovadas (dados tratados da gerenciadora)
    if not df_aprov_ger.empty:
        cols_aprov = [c for c in df_aprov_ger.columns if c != "cliente_os"]
        _aba_dados(wb, df_aprov_ger[cols_aprov].reset_index(drop=True), "Aprovadas")


# =============================================================
#  MAIN
# =============================================================

def main():
    print("=" * 60)
    print("  Processador OS + APROVADAS por Gerenciadora")
    print("=" * 60)

    # ── Arquivo de OS ────────────────────────────────────────────
    arquivo_os = escolher_arquivo()
    if not arquivo_os:
        input("\nPressione Enter para sair...")
        return
    print(f"\nOS:        {os.path.basename(arquivo_os)}")

    # ── Trata OS ─────────────────────────────────────────────────
    try:
        df_os = tratar_dados_os(arquivo_os)
    except Exception as e:
        print(f"\nERRO ao tratar OS: {e}")
        input("\nPressione Enter para sair...")
        return

    # ── Trata APROVADAS ──────────────────────────────────────────
    # A coluna Gerenciadora já vem da própria planilha APROVADAS
    try:
        df_aprov = tratar_dados_aprovadas()
        if not df_aprov.empty:
            print(f"APROVADAS: {os.path.basename(CONFIG['aprovadas_file'])}")
    except Exception as e:
        print(f"\nERRO ao tratar APROVADAS: {e}")
        df_aprov = pd.DataFrame()

    # ── Base bruta OS por gerenciadora (sem filtro 60d) ──────────
    tabela_carteira = pd.read_excel(CONFIG["carteiras_file"], sheet_name=0)
    tabela_carteira.columns = tabela_carteira.columns.str.strip()
    cart_base = ["Cod"]
    if "Gerenciadora" in tabela_carteira.columns:
        cart_base.append("Gerenciadora")

    df_raw_os = pd.read_excel(arquivo_os, sheet_name=0)
    df_raw_os.columns = df_raw_os.columns.str.strip()
    df_raw_os = pd.merge(df_raw_os, tabela_carteira[cart_base],
                         left_on="Cod Cliente", right_on="Cod", how="left")
    df_raw_os.drop(columns="Cod", inplace=True, errors="ignore")
    df_raw_os = df_raw_os[~df_raw_os["Status OS"].isin(STATUS_INDESEJADOS)].copy()
    if "Cod Cliente" in df_raw_os.columns and "Cod OS" in df_raw_os.columns:
        df_raw_os["_k"] = (df_raw_os["Cod Cliente"].astype(str) + "-" +
                           df_raw_os["Cod OS"].astype(str))
        df_raw_os = df_raw_os.drop_duplicates(subset="_k")

    # ── Base bruta APROVADAS por gerenciadora (sem filtro 60d) ───
    # A coluna Gerenciadora já está na planilha — só agrupa diretamente
    base_aprov_por_ger = {}
    if not df_aprov.empty and os.path.exists(CONFIG["aprovadas_file"]):
        try:
            df_raw_ap = pd.read_excel(CONFIG["aprovadas_file"], sheet_name=0)
            df_raw_ap.columns = df_raw_ap.columns.str.strip()
            df_raw_ap["_rec"] = (df_raw_ap["Valor Total de Itens"].fillna(0) +
                                 df_raw_ap["Valor Total de MDO"].fillna(0))
            # Agrupa por Gerenciadora (coluna já existente na planilha)
            for ger_k, grp in df_raw_ap.groupby("Gerenciadora"):
                base_aprov_por_ger[str(ger_k).strip()] = {
                    "qtd":     int(len(grp)),
                    "receita": float(grp["_rec"].sum()),
                }
        except Exception as e:
            log.warning(f"Nao foi possivel calcular base bruta de APROVADAS: {e}")

    # ── Lista de gerenciadoras — union de OS e APROVADAS ─────────
    gers_os    = set(df_os["Gerenciadora"].dropna().unique()) if "Gerenciadora" in df_os.columns else set()
    gers_aprov = set(df_aprov["Gerenciadora"].dropna().unique()) if not df_aprov.empty and "Gerenciadora" in df_aprov.columns else set()
    geradoras  = sorted(gers_os | gers_aprov) or ["Geral"]

    os.makedirs(CONFIG["pasta_saida"], exist_ok=True)
    ts = dt.datetime.now().strftime("%Y%m%d_%H%M%S")

    gerados = []
    for ger in geradoras:
        ger_upper = str(ger).strip().upper()

        # Dados OS desta gerenciadora
        if "Gerenciadora" not in df_os.columns:
            df_os_ger = df_os.copy()
        else:
            df_os_ger = df_os[df_os["Gerenciadora"].str.strip().str.upper() == ger_upper].copy()

        # Base bruta OS desta gerenciadora
        if "Gerenciadora" not in df_raw_os.columns:
            base_os_ger = len(df_raw_os)
        else:
            base_os_ger = int(
                (df_raw_os["Gerenciadora"].str.strip().str.upper() == ger_upper).sum()
            )

        # Dados APROVADAS desta gerenciadora
        if df_aprov.empty or "Gerenciadora" not in df_aprov.columns:
            df_aprov_ger = pd.DataFrame()
        else:
            df_aprov_ger = df_aprov[
                df_aprov["Gerenciadora"].str.strip().str.upper() == ger_upper
            ].copy()

        # Base bruta APROVADAS desta gerenciadora
        # Busca pelo nome exato (case-insensitive) no dicionário
        base_aprov_ger = next(
            (v for k, v in base_aprov_por_ger.items() if k.upper() == ger_upper),
            {"qtd": 0, "receita": 0.0},
        )

        # Históricos separados por gerenciadora e fonte
        hist_os    = carregar_historico(ger, "os")
        hist_aprov = carregar_historico(ger, "aprov")

        salvar_historico_os(df_os_ger, ger, hist_os)
        if not df_aprov_ger.empty:
            salvar_historico_aprov(df_aprov_ger, ger, hist_aprov)

        # Gera o workbook
        nome = f"OS_{ger}_{ts}.xlsx".replace("/", "-").replace(" ", "_")
        path = os.path.join(CONFIG["pasta_saida"], nome)
        wb   = Workbook()
        wb.remove(wb.active)
        build_sheet(
            wb, ger,
            df_os_ger,    hist_os,    base_os_ger,
            df_aprov_ger, hist_aprov, base_aprov_ger,
        )
        wb.save(path)
        gerados.append(path)
        print(f"  Gerado: {nome}  (OS={len(df_os_ger)} | Aprov={len(df_aprov_ger)})")

    print(f"\n✅ {len(gerados)} relatorio(s) salvo(s) em:")
    print(f"   {CONFIG['pasta_saida']}")
    input("\nPressione Enter para sair...")


if __name__ == "__main__":
    main()
